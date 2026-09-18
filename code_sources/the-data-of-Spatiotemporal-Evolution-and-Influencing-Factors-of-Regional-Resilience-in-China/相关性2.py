# -*- coding: utf-8 -*-
"""
计算解释变量与四个区域韧性子系统之间的相关性。

子系统变量：
- ER  ：Economic Resilience
- SDR ：Social Development Resilience
- EER ：Ecological and Environmental Resilience
- IR  ：Infrastructure Resilience

解释变量：
- URB, GOV, INNO, EDU, OPENPC
- FIN, DIG, PRI, INC, SOC

输出文件：
CORRELATION_RESULTS/
    TABLE_A3_SUBSYSTEM_CORRELATION.xlsx

Excel工作表：
1. Table_A3     带显著性星号的论文表格
2. Coefficients 相关系数
3. P_Values     p值
4. Sample_Size  有效样本量
"""

from pathlib import Path
from typing import Optional, Tuple

import numpy as np
import pandas as pd
from scipy.stats import pearsonr, spearmanr
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


# =========================================================
# 1. File settings
# =========================================================
INPUT_PATH = Path(
    r"PANEL_ANALYSIS_DATA - 副本 - 副本.xlsx"
)

# None表示自动查找包含全部所需变量的工作表
# 如果已经知道工作表名称，可修改为：
# INPUT_SHEET = "PANEL_ANALYSIS"
INPUT_SHEET: Optional[str] = None

OUTPUT_DIR = Path("CORRELATION_RESULTS")
OUTPUT_PATH = (
    OUTPUT_DIR
    / "TABLE_A3_SUBSYSTEM_CORRELATION.xlsx"
)


# =========================================================
# 2. Variable settings
# =========================================================
ID_COLS = [
    "ADMINCODE",
    "ZONE",
    "REGION",
    "YEAR",
]

SUBSYSTEM_COLS = [
    "ER",
    "SDR",
    "EER",
    "IR",
]

EXPLANATORY_COLS = [
    "URB",
    "GOV",
    "INNO",
    "EDU",
    "OPENPC",
    "FIN",
    "DIG",
    "PRI",
    "INC",
    "SOC",
]

# 与原Figure 4保持一致，默认使用Pearson相关系数
# 如需Spearman相关系数，将其改为：
# CORRELATION_METHOD = "spearman"
CORRELATION_METHOD = "pearson"

# 显示的小数位数
DECIMAL_PLACES = 3


# =========================================================
# 3. Utility functions
# =========================================================
def clean_colname(name: str) -> str:
    """Clean spaces and special spaces in column names."""
    return (
        str(name)
        .strip()
        .replace("\u3000", " ")
        .replace("\n", "")
        .replace("\r", "")
    )


def safe_to_numeric(series: pd.Series) -> pd.Series:
    """Convert a series to numeric and remove commas."""
    return pd.to_numeric(
        series.astype(str)
        .str.replace(",", "", regex=False)
        .str.strip(),
        errors="coerce",
    )


def significance_stars(p_value: float) -> str:
    """
    Return significance symbols.

    *** p < 0.001
    **  p < 0.01
    *   p < 0.05
    """
    if pd.isna(p_value):
        return ""

    if p_value < 0.001:
        return "***"
    if p_value < 0.01:
        return "**"
    if p_value < 0.05:
        return "*"

    return ""


# =========================================================
# 4. Read the Excel file
# =========================================================
def find_valid_sheet(
    excel_file: pd.ExcelFile,
) -> str:
    """
    Automatically find a worksheet containing:
    ER, SDR, EER, IR and all explanatory variables.
    """
    required_cols = set(
        SUBSYSTEM_COLS + EXPLANATORY_COLS
    )

    sheet_information = []

    for sheet_name in excel_file.sheet_names:
        preview = pd.read_excel(
            excel_file,
            sheet_name=sheet_name,
            nrows=5,
        )

        preview.columns = [
            clean_colname(col)
            for col in preview.columns
        ]

        available_cols = set(preview.columns)
        missing_cols = sorted(
            required_cols - available_cols
        )

        sheet_information.append(
            {
                "sheet": sheet_name,
                "missing": missing_cols,
            }
        )

        if not missing_cols:
            return sheet_name

    details = "\n".join(
        [
            f"- {item['sheet']}: "
            f"缺少 {item['missing']}"
            for item in sheet_information
        ]
    )

    raise ValueError(
        "未找到同时包含全部解释变量和四个子系统"
        "韧性得分的工作表。\n\n"
        "各工作表检查结果：\n"
        f"{details}"
    )


def read_data(
    input_path: Path,
    input_sheet: Optional[str],
) -> Tuple[pd.DataFrame, str]:
    """Read Excel or CSV data."""
    if not input_path.exists():
        raise FileNotFoundError(
            f"未找到输入文件：{input_path.resolve()}"
        )

    suffix = input_path.suffix.lower()

    if suffix in [".xlsx", ".xls"]:
        excel_file = pd.ExcelFile(input_path)

        if input_sheet is None:
            selected_sheet = find_valid_sheet(
                excel_file
            )
        else:
            if input_sheet not in excel_file.sheet_names:
                raise ValueError(
                    f"工作表不存在：{input_sheet}\n"
                    f"现有工作表：{excel_file.sheet_names}"
                )

            selected_sheet = input_sheet

        df = pd.read_excel(
            excel_file,
            sheet_name=selected_sheet,
        )

    elif suffix == ".csv":
        selected_sheet = "CSV"
        df = pd.read_csv(
            input_path,
            encoding="utf-8-sig",
        )

    else:
        raise ValueError(
            "仅支持.xlsx、.xls或.csv文件。"
        )

    df.columns = [
        clean_colname(col)
        for col in df.columns
    ]

    return df, selected_sheet


# =========================================================
# 5. Data preparation
# =========================================================
def prepare_data(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """Check columns and convert variables to numeric."""
    required_cols = (
        SUBSYSTEM_COLS
        + EXPLANATORY_COLS
    )

    missing_cols = [
        col
        for col in required_cols
        if col not in df.columns
    ]

    if missing_cols:
        raise ValueError(
            "数据缺少以下必要字段：\n"
            + "\n".join(missing_cols)
        )

    keep_cols = [
        col
        for col in ID_COLS
        if col in df.columns
    ] + required_cols

    result = df[keep_cols].copy()

    if "YEAR" in result.columns:
        result["YEAR"] = safe_to_numeric(
            result["YEAR"]
        )

    for col in required_cols:
        result[col] = safe_to_numeric(
            result[col]
        )

        # 将正负无穷值视为缺失值
        result[col] = result[col].replace(
            [np.inf, -np.inf],
            np.nan,
        )

    # 删除四个子系统和全部解释变量均为空的行
    result = result.dropna(
        subset=required_cols,
        how="all",
    ).reset_index(drop=True)

    return result


# =========================================================
# 6. Correlation calculation
# =========================================================
def calculate_single_correlation(
    x: pd.Series,
    y: pd.Series,
    method: str,
) -> Tuple[float, float, int]:
    """
    Calculate correlation using pairwise complete samples.

    Returns:
    correlation coefficient, p-value, sample size
    """
    pair = pd.DataFrame(
        {
            "x": x,
            "y": y,
        }
    ).dropna()

    n = len(pair)

    if n < 3:
        return np.nan, np.nan, n

    # Correlation cannot be calculated for constants
    if (
        pair["x"].nunique() < 2
        or pair["y"].nunique() < 2
    ):
        return np.nan, np.nan, n

    if method.lower() == "pearson":
        coefficient, p_value = pearsonr(
            pair["x"].to_numpy(),
            pair["y"].to_numpy(),
        )

    elif method.lower() == "spearman":
        coefficient, p_value = spearmanr(
            pair["x"].to_numpy(),
            pair["y"].to_numpy(),
        )

    else:
        raise ValueError(
            "CORRELATION_METHOD只能设置为"
            "'pearson'或'spearman'。"
        )

    return (
        float(coefficient),
        float(p_value),
        int(n),
    )


def calculate_correlations(
    df: pd.DataFrame,
) -> Tuple[
    pd.DataFrame,
    pd.DataFrame,
    pd.DataFrame,
]:
    """
    Calculate correlations between explanatory variables
    and ER, SDR, EER and IR.
    """
    coefficient_df = pd.DataFrame(
        index=EXPLANATORY_COLS,
        columns=SUBSYSTEM_COLS,
        dtype=float,
    )

    pvalue_df = pd.DataFrame(
        index=EXPLANATORY_COLS,
        columns=SUBSYSTEM_COLS,
        dtype=float,
    )

    sample_size_df = pd.DataFrame(
        index=EXPLANATORY_COLS,
        columns=SUBSYSTEM_COLS,
        dtype=float,
    )

    for feature in EXPLANATORY_COLS:
        for subsystem in SUBSYSTEM_COLS:
            coefficient, p_value, n = (
                calculate_single_correlation(
                    df[feature],
                    df[subsystem],
                    CORRELATION_METHOD,
                )
            )

            coefficient_df.loc[
                feature,
                subsystem,
            ] = coefficient

            pvalue_df.loc[
                feature,
                subsystem,
            ] = p_value

            sample_size_df.loc[
                feature,
                subsystem,
            ] = n

    sample_size_df = sample_size_df.astype(
        "Int64"
    )

    return (
        coefficient_df,
        pvalue_df,
        sample_size_df,
    )


# =========================================================
# 7. Build publication table
# =========================================================
def build_formatted_table(
    coefficient_df: pd.DataFrame,
    pvalue_df: pd.DataFrame,
) -> pd.DataFrame:
    """Create coefficient table with significance stars."""
    table = pd.DataFrame(
        index=EXPLANATORY_COLS,
        columns=SUBSYSTEM_COLS,
        dtype=object,
    )

    for feature in EXPLANATORY_COLS:
        for subsystem in SUBSYSTEM_COLS:
            coefficient = coefficient_df.loc[
                feature,
                subsystem,
            ]

            p_value = pvalue_df.loc[
                feature,
                subsystem,
            ]

            if pd.isna(coefficient):
                table.loc[
                    feature,
                    subsystem,
                ] = ""
                continue

            stars = significance_stars(p_value)

            table.loc[
                feature,
                subsystem,
            ] = (
                f"{coefficient:.{DECIMAL_PLACES}f}"
                f"{stars}"
            )

    table.index.name = "Variable"

    return table.reset_index()


# =========================================================
# 8. Excel formatting
# =========================================================
def format_excel(
    output_path: Path,
) -> None:
    """Apply simple academic-table formatting."""
    from openpyxl import load_workbook

    workbook = load_workbook(output_path)

    thin_black = Side(
        style="thin",
        color="000000",
    )

    medium_black = Side(
        style="medium",
        color="000000",
    )

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "B2"

        max_row = worksheet.max_row
        max_col = worksheet.max_column

        # Header formatting
        for cell in worksheet[1]:
            cell.font = Font(
                name="Times New Roman",
                size=11,
                bold=True,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

            cell.border = Border(
                top=medium_black,
                bottom=thin_black,
            )

        # Body formatting
        for row in worksheet.iter_rows(
            min_row=2,
            max_row=max_row,
            min_col=1,
            max_col=max_col,
        ):
            for cell in row:
                cell.font = Font(
                    name="Times New Roman",
                    size=11,
                )

                if cell.column == 1:
                    cell.alignment = Alignment(
                        horizontal="left",
                        vertical="center",
                    )
                else:
                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                    )

        # Bottom border
        for cell in worksheet[max_row]:
            cell.border = Border(
                bottom=medium_black
            )

        # Column width
        for column_index in range(
            1,
            max_col + 1,
        ):
            letter = get_column_letter(
                column_index
            )

            if column_index == 1:
                worksheet.column_dimensions[
                    letter
                ].width = 18
            else:
                worksheet.column_dimensions[
                    letter
                ].width = 16

        worksheet.row_dimensions[1].height = 23

    workbook.save(output_path)


# =========================================================
# 9. Save results
# =========================================================
def save_results(
    formatted_table: pd.DataFrame,
    coefficient_df: pd.DataFrame,
    pvalue_df: pd.DataFrame,
    sample_size_df: pd.DataFrame,
) -> None:
    """Save all results into one Excel workbook."""
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    coefficient_output = (
        coefficient_df
        .round(6)
        .reset_index()
        .rename(columns={"index": "Variable"})
    )

    pvalue_output = (
        pvalue_df
        .round(6)
        .reset_index()
        .rename(columns={"index": "Variable"})
    )

    sample_size_output = (
        sample_size_df
        .reset_index()
        .rename(columns={"index": "Variable"})
    )

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl",
    ) as writer:
        formatted_table.to_excel(
            writer,
            sheet_name="Table_A3",
            index=False,
        )

        coefficient_output.to_excel(
            writer,
            sheet_name="Coefficients",
            index=False,
        )

        pvalue_output.to_excel(
            writer,
            sheet_name="P_Values",
            index=False,
        )

        sample_size_output.to_excel(
            writer,
            sheet_name="Sample_Size",
            index=False,
        )

    format_excel(OUTPUT_PATH)


# =========================================================
# 10. Main
# =========================================================
def main() -> None:
    print("=" * 72)
    print(
        "Correlation analysis between explanatory "
        "variables and resilience subsystems"
    )
    print("=" * 72)

    df, selected_sheet = read_data(
        INPUT_PATH,
        INPUT_SHEET,
    )

    print(f"输入文件：{INPUT_PATH.resolve()}")
    print(f"使用工作表：{selected_sheet}")

    df = prepare_data(df)

    print(f"有效数据行数：{len(df)}")
    print(
        "相关性方法："
        f"{CORRELATION_METHOD.capitalize()}"
    )

    (
        coefficient_df,
        pvalue_df,
        sample_size_df,
    ) = calculate_correlations(df)

    formatted_table = build_formatted_table(
        coefficient_df,
        pvalue_df,
    )

    save_results(
        formatted_table,
        coefficient_df,
        pvalue_df,
        sample_size_df,
    )

    print("\n带显著性标记的结果：")
    print(
        formatted_table.to_string(
            index=False
        )
    )

    print("\n显著性说明：")
    print("* p < 0.05")
    print("** p < 0.01")
    print("*** p < 0.001")

    print(
        "\n输出文件："
        f"{OUTPUT_PATH.resolve()}"
    )

    print("\n输出工作表：")
    print("1. Table_A3：论文使用的相关性表格")
    print("2. Coefficients：相关系数")
    print("3. P_Values：显著性p值")
    print("4. Sample_Size：配对有效样本量")
    print("=" * 72)


if __name__ == "__main__":
    main()